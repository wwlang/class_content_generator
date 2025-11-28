"""
Big number slide handler.

Handles slides with large centered statistic/number and explanation text.
"""

from typing import Any
from lxml import etree
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Pt

from .base import SlideHandler
from ..renderers import BaseRenderer
from ..config import LayoutConfig, FontConfig, ColorConfig


class BigNumberSlideHandler(SlideHandler):
    """
    Handler for big number slides.

    Big number slides contain:
    - Large centered number (135pt, orange)
    - Explanation text below (medium size, muted gray)
    - Cream background
    """

    @property
    def priority(self) -> int:
        """Big number slides have high priority (checked early)."""
        return 20

    def can_handle(self, html_slide: etree.Element) -> bool:
        """
        Check if this is a big number slide.

        Args:
            html_slide: HTML slide element

        Returns:
            True if slide has 'big-number-slide' class
        """
        return self._has_class(html_slide, 'big-number-slide')

    def handle(self, slide: Any, html_slide: etree.Element) -> None:
        """
        Process big number slide using renderers.

        Args:
            slide: PowerPoint slide to populate
            html_slide: HTML slide element with content
        """
        # Initialize renderer
        base_renderer = BaseRenderer(slide)

        # Check for dark background
        is_dark = self._has_class(html_slide, 'dark-bg')

        # Apply background
        if is_dark:
            base_renderer.apply_background(ColorConfig.DARK_GRAY)
        else:
            base_renderer.apply_background(ColorConfig.CREAM)

        # Find number content container
        number_content = html_slide.find('.//*[@class="number-content"]')
        if number_content is None:
            return

        # Extract and render big number
        big_number_elem = number_content.find('.//*[@class="big-number"]')
        if big_number_elem is not None:
            number_text = self.converter.extract_text_content(big_number_elem).strip()

            # Create centered text box for big number
            number_box = base_renderer.add_textbox(
                1.0,  # Centered horizontally with padding
                LayoutConfig.BIG_NUMBER_Y,
                LayoutConfig.SLIDE_WIDTH - 2.0,
                LayoutConfig.BIG_NUMBER_HEIGHT
            )
            tf = number_box.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER

            run = p.add_run()
            run.text = number_text
            run.font.name = FontConfig.HEADER_FONT
            run.font.size = FontConfig.BIG_NUMBER_SIZE
            run.font.bold = True
            run.font.color.rgb = ColorConfig.ORANGE

        # Extract and render explanation text
        explanation_elem = number_content.find('.//*[@class="number-explanation"]')
        if explanation_elem is not None:
            # Create text box for explanation below number
            explanation_box = base_renderer.add_textbox(
                1.0,
                LayoutConfig.BIG_NUMBER_Y + LayoutConfig.BIG_NUMBER_EXPLANATION_Y_OFFSET,
                LayoutConfig.SLIDE_WIDTH - 2.0,
                LayoutConfig.BIG_NUMBER_EXPLANATION_HEIGHT
            )
            tf = explanation_box.text_frame
            tf.word_wrap = True

            # Set default text color based on background
            text_color = ColorConfig.CREAM if is_dark else ColorConfig.MUTED_GRAY

            # Process all <p> elements in the explanation
            p_elements = explanation_elem.findall('.//p')

            if p_elements:
                # First paragraph
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                self.converter.add_formatted_text(p, p_elements[0])

                # Apply formatting to runs
                for run in p.runs:
                    if not run.font.name:
                        run.font.name = FontConfig.BODY_FONT
                    if not run.font.size:
                        run.font.size = FontConfig.BODY_MEDIUM
                    # Only set color if not bold (bold text stays orange)
                    if not run.font.bold:
                        run.font.color.rgb = text_color

                # Additional paragraphs
                for p_elem in p_elements[1:]:
                    p_new = tf.add_paragraph()
                    p_new.alignment = PP_ALIGN.CENTER
                    p_new.space_before = Pt(6)  # Small space between paragraphs
                    self.converter.add_formatted_text(p_new, p_elem)

                    # Apply formatting to runs
                    for run in p_new.runs:
                        if not run.font.name:
                            run.font.name = FontConfig.BODY_FONT
                        if not run.font.size:
                            run.font.size = FontConfig.BODY_MEDIUM
                        if not run.font.bold:
                            run.font.color.rgb = text_color
            else:
                # Fallback: no <p> tags, use entire element
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                self.converter.add_formatted_text(p, explanation_elem)

                for run in p.runs:
                    if not run.font.name:
                        run.font.name = FontConfig.BODY_FONT
                    if not run.font.size:
                        run.font.size = FontConfig.BODY_MEDIUM
                    if not run.font.bold:
                        run.font.color.rgb = text_color
